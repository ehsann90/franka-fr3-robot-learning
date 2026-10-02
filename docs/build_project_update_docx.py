#!/usr/bin/env python3
"""Build the editable source for the FR3 research progress report PDF."""

from __future__ import annotations

from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "FR3_PROJECT_UPDATE_FOR_HUYUE_AND_DR_MEI.html"
OUTPUT = HERE / "FR3_PROJECT_UPDATE_FOR_HUYUE_AND_DR_MEI.docx"


class ContentParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.kind: str | None = None
        self.parts: list[str] = []
        self.blocks: list[tuple[str, str]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"h1", "h2", "p"}:
            self.kind = tag
            self.parts = []

    def handle_data(self, data: str) -> None:
        if self.kind is not None:
            self.parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == self.kind:
            self.blocks.append((tag, "".join(self.parts).strip()))
            self.kind = None
            self.parts = []


def main() -> None:
    parser = ContentParser()
    parser.feed(SOURCE.read_text(encoding="utf-8"))
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(5)
    for kind, content in parser.blocks:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.keep_with_next = kind in {"h1", "h2"}
        if kind == "h1":
            paragraph.paragraph_format.space_after = Pt(10)
        elif kind == "h2":
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(content)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.bold = kind in {"h1", "h2"}
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
